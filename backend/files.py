"""Save uploaded company files and extract their text synchronously.

Supported: .txt, .md, .pdf, .docx, .csv, .xlsx. Max 10 MB. Extraction runs on
upload; if it fails the file is still stored with status 'extract_failed' so the
UI can warn and recommendation/chat can continue without it.
"""

import csv
import io
import logging
from pathlib import Path

import docx
import openpyxl
from pypdf import PdfReader

from backend import db

log = logging.getLogger(__name__)

MAX_FILE_BYTES = 10 * 1024 * 1024  # 10 MB
ALLOWED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".csv", ".xlsx"}

# Per-file storage cap; keeps one huge file from filling the DB with text no
# prompt could use anyway (max_grant file size is 10 MB, extracted text could
# otherwise be far larger for a text-heavy format).
MAX_STORED_TEXT_CHARS = 80_000

# Total combined budget for company-file context sent to the LLM (/recommend
# and measure chat share this limit, per the plan's context rules).
MAX_CONTEXT_CHARS = 80_000


def is_allowed(filename: str) -> bool:
    """True if the filename's extension is one we know how to extract."""
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS


def safe_filename(filename: str) -> str:
    """Strip every directory component from a client-supplied filename.

    The name arrives verbatim in the multipart part and is never trustworthy:
    '../../../x.txt' would otherwise land outside the vestlus's upload folder.
    Backslashes are folded to '/' first — they are legal POSIX filename
    characters, so Path().name would keep a Windows-style path as one very
    confusing single name instead of reducing it to its last segment.
    """
    name = Path(filename.replace("\\", "/")).name.strip()
    if name in ("", ".", "..") or "\x00" in name:
        raise ValueError(f"unusable filename: {filename!r}")
    return name


def _extract_txt(data: bytes) -> str:
    """Decode plain text, replacing any invalid bytes rather than raising."""
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    """Extract text from each page of a PDF and join with newlines."""
    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    """Extract paragraph text from a Word document."""
    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)


def _extract_csv(data: bytes) -> str:
    """Render CSV rows as comma-joined lines of plain text."""
    text = data.decode("utf-8", errors="replace")
    rows = csv.reader(io.StringIO(text))
    return "\n".join(", ".join(row) for row in rows)


def _extract_xlsx(data: bytes) -> str:
    """Render every sheet's cells as plain text, one sheet header + rows each."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            parts.append(", ".join(cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from file bytes. Raises on failure."""
    ext = Path(filename).suffix.lower()
    if ext in (".txt", ".md"):
        text = _extract_txt(data)
    elif ext == ".pdf":
        text = _extract_pdf(data)
    elif ext == ".docx":
        text = _extract_docx(data)
    elif ext == ".csv":
        text = _extract_csv(data)
    elif ext == ".xlsx":
        text = _extract_xlsx(data)
    else:
        raise ValueError(f"unsupported extension: {ext}")

    # Cap from the front so we keep the most recent/later content in the file.
    if len(text) > MAX_STORED_TEXT_CHARS:
        text = text[-MAX_STORED_TEXT_CHARS:]
    return text


def save_upload(vestluse_id: str, filename: str, data: bytes) -> dict:
    """Store bytes on disk and a DB row with extracted text (or extract_failed).

    Sanitises the name itself rather than trusting the caller — this is the one
    place a client-supplied string becomes a filesystem path.
    """
    filename = safe_filename(filename)
    upload_dir = db.UPLOADS_DIR / vestluse_id
    upload_dir.mkdir(parents=True, exist_ok=True)
    dest = upload_dir / filename
    dest.write_bytes(data)

    try:
        text = extract_text(filename, data)
        status = "ok"
    except Exception as e:  # extraction is best-effort; store the file regardless
        log.warning("extraction failed for %s: %s", filename, e)
        text = None
        status = "extract_failed"

    file_id = db.add_file(vestluse_id, filename, str(dest), text, status)
    return {"id": file_id, "name": filename, "status": status}


def build_context_text(vestluse_id: str) -> str:
    """Join every ok-status file's extracted text into one prompt block.

    Used by /recommend and measure chat so both see the same company context.
    Files are joined oldest-to-newest, then the combined text is trimmed from
    the front (dropping the oldest content) if it exceeds MAX_CONTEXT_CHARS.
    """
    entries = db.get_extracted_texts(vestluse_id)
    if not entries:
        return ""
    blocks = [f"## Fail: {e['name']}\n{e['extracted_text']}" for e in entries]
    combined = "\n\n".join(blocks)
    if len(combined) > MAX_CONTEXT_CHARS:
        combined = combined[-MAX_CONTEXT_CHARS:]
    return combined
